"""Import du registre GHPL (classement général des scans) : complète sans écraser."""

from datetime import date

import pytest
from openpyxl import Workbook

from apps.amm.models import MarketingAuthorization, Renewal
from apps.catalog.models import Product
from apps.imports.models import ImportBatch, ImportRow
from apps.imports.registry import is_registry
from apps.imports.services import import_workbook

CONSOLIDATION = [
    "Pays",
    "Gamme",
    "Produit",
    "Presentation",
    "Source",
    "Controle",
    "Statut dossier",
    "Referentiel",
    "N° AMM origine",
    "Date origine",
]
DOCS = ["Pays", "Gamme", "Produit", "Presentation", "Type", "Date", "N° AMM"]


def registry(path, records, docs=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "_Consolidation"
    sheet.append(CONSOLIDATION)
    for record in records:
        sheet.append([record.get(column) for column in CONSOLIDATION])
    other = workbook.create_sheet("_Docs")
    other.append(DOCS)
    for doc in docs:
        other.append([doc.get(column) for column in DOCS])
    workbook.create_sheet("Parametres pays")
    target = path / "0_REGISTRE AMM GHPL.xlsx"
    workbook.save(target)
    return str(target)


def row(**values):
    return {
        "Pays": "SENEGAL",
        "Gamme": "GAMME GENERALE",
        "Source": "Dashboard + classement",
        "Controle": "Concordant",
        **values,
    }


@pytest.fixture
def batch(users):
    return ImportBatch.objects.create(created_by=users["ceo"])


def test_registry_detected(tmp_path, workbook_path):
    assert is_registry(registry(tmp_path, []))
    assert not is_registry(str(workbook_path))


def test_completes_without_overwriting(tmp_path, make_amm, ranges, batch):
    empty = make_amm(
        "SN",
        product_obj=Product.objects.create(name="LOLIP 10MG CPR B/30"),
        original_number="",
        original_start_date=None,
    )
    filled = make_amm(
        "SN",
        product_obj=Product.objects.create(name="GENOPRIL 10MG CPR B/30"),
        original_number="AMM-DASHBOARD",
    )
    source = registry(
        tmp_path,
        [
            row(
                Presentation="LOLIP 10MG CPR B/30",
                **{"N° AMM origine": "19/0042", "Date origine": date(2019, 3, 4)},
            ),
            row(Presentation="GENOPRIL 10MG CPR B/30", **{"N° AMM origine": "AMM-REGISTRE"}),
        ],
    )
    summary = import_workbook(source, batch=batch)
    assert summary["kind"] == "registry"
    assert summary["sheets"]["Registre GHPL — SENEGAL"]["updated"] == 1
    empty.refresh_from_db()
    filled.refresh_from_db()
    assert empty.original_number == "19/0042"
    assert empty.original_start_date == date(2019, 3, 4)
    assert filled.original_number == "AMM-DASHBOARD"
    message = ImportRow.objects.get(batch=batch, outcome="UPDATED").message
    assert message.startswith("Complété : n° d'origine 19/0042")
    assert empty.history.first().history_change_reason.startswith("Complété depuis le registre")


def test_scans_only_presentation_reuses_catalog_product(
    tmp_path, countries, ranges, make_amm, batch
):
    """« OMEPRAL 20MG GELULE B28 » est « OMEPRAL 20MG GEL B/28 », suivi au Mali."""
    omepral = Product.objects.create(name="OMEPRAL 20MG GEL B/28", range=ranges["GENERALE"])
    make_amm("ML", product_obj=omepral)
    tenso = Product.objects.create(name="TENSOPLUS 10MG/2,5MG/10MG CPR B/30")
    make_amm("ML", product_obj=tenso)
    source = registry(
        tmp_path,
        [
            row(
                Presentation="OMEPRAL 20MG GELULE B28",
                Source="Classement seul",
                **{"N° AMM origine": "SN-7"},
            ),
            row(Presentation="TENSOPLUS 2,5MG-10MG-10MG CPR B30", Source="Classement seul"),
            row(Presentation="DICAGEN 1G PDRE INJ", Source="Classement seul", Gamme="GAMME CARDIO"),
            row(Presentation="AMLOPERIN - PRESENTATION NON IDENTIFIEE", Source="Classement seul"),
        ],
    )
    import_workbook(source, batch=batch)
    senegal = countries["SN"]
    amm = MarketingAuthorization.objects.get(product=omepral, country=senegal)
    assert amm.original_number == "SN-7"
    assert "absente du Dashboard" in amm.notes
    assert MarketingAuthorization.objects.filter(product=tenso, country=senegal).exists()
    created = Product.objects.get(name="DICAGEN 1G PDRE INJ")
    assert created.range.code == "CARDIO"
    assert not Product.objects.filter(name__contains="AMLOPERIN").exists()
    assert Product.objects.filter(name__startswith="OMEPRAL").count() == 1
    assert Product.objects.filter(name__startswith="TENSOPLUS").count() == 1
    messages = dict(ImportRow.objects.filter(batch=batch).values_list("row_number", "message"))
    assert messages[2].startswith("Produit du catalogue repris : OMEPRAL 20MG GEL B/28")
    assert messages[4].startswith("Produit créé")
    assert messages[5].startswith("Présentation non identifiée")


def test_deposited_opens_renewal_and_flags_newer_acts(tmp_path, make_amm, make_renewal, batch):
    product = Product.objects.create(name="ALFA GH 10MG LP CPR B30")
    amm = make_amm("CI", product_obj=product)
    other = make_amm("CI", product_obj=Product.objects.create(name="BISOPROLOL GH 5MG CPR B/30"))
    make_renewal(other, status="EN_INSTRUCTION")
    source = registry(
        tmp_path,
        [
            row(
                Pays="COTE D'IVOIRE",
                Presentation="ALFA-GH 10MG CPR B30",
                Referentiel="ALFA GH 10MG LP CPR B30",
                **{"Statut dossier": "Depose (attestation)"},
            ),
            row(
                Pays="COTE D'IVOIRE",
                Presentation="BISOPROLOL GH 5MG CPR B/30",
                Controle="acte classe plus recent que le Dashboard (2026-01-13) - a verifier",
                **{"Statut dossier": "Depose (attestation)"},
            ),
            row(Pays="RDC", Presentation="LOLIP 10MG CPR B/30"),
            row(Presentation="INCONNU 5MG CPR B/30"),
        ],
        docs=[
            {
                "Pays": "COTE D'IVOIRE",
                "Presentation": "ALFA GH 10MG LP CPR B30",
                "Type": "Depot / attestation",
                "Date": date(2025, 11, 6),
            },
            {
                "Pays": "COTE D'IVOIRE",
                "Presentation": "ALFA GH 10MG LP CPR B30",
                "Type": "Depot / attestation",
                "Date": date(2024, 2, 1),
            },
        ],
    )
    import_workbook(source, batch=batch)
    renewal = amm.renewals.get()
    assert renewal.workflow_status == Renewal.WorkflowStatus.DEPOSE
    assert renewal.filing_date == date(2025, 11, 6)
    assert other.renewals.count() == 1  # déjà en instruction : rien d'ouvert en double
    rows = {r.row_number: r for r in ImportRow.objects.filter(batch=batch)}
    assert rows[2].outcome == "UPDATED"
    assert "attestation du 06/11/2025" in rows[2].message
    assert rows[3].outcome == "WARNING"
    assert "acte classé du 13/01/2026 est plus récent" in rows[3].message
    assert rows[4].outcome == "WARNING" and "Pays non suivi" in rows[4].message
    assert rows[5].outcome == "WARNING" and "AMM du Dashboard introuvable" in rows[5].message


def test_dry_run_changes_nothing(tmp_path, make_amm, batch):
    amm = make_amm(
        "SN",
        product_obj=Product.objects.create(name="LOLIP 10MG CPR B/30"),
        original_number="",
    )
    source = registry(
        tmp_path,
        [
            row(
                Presentation="LOLIP 10MG CPR B/30",
                **{"N° AMM origine": "19/0042", "Statut dossier": "Depose"},
            ),
            row(Presentation="NOUVEAU 5MG CPR B/30", Source="Classement seul"),
        ],
    )
    summary = import_workbook(source, batch=batch, dry_run=True)
    assert summary["totals"]["created"] == 1 and summary["totals"]["updated"] == 1
    amm.refresh_from_db()
    assert amm.original_number == ""
    assert not amm.renewals.exists()
    assert not Product.objects.filter(name="NOUVEAU 5MG CPR B/30").exists()
    assert ImportRow.objects.filter(batch=batch, amm__isnull=True).count() == 2
